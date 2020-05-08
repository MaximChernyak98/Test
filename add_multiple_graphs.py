import openpyxl
import datetime


from openpyxl import load_workbook
from openpyxl.styles import Alignment
from datetime import datetime, date, time

from openpyxl.chart import (LineChart, Reference, ScatterChart, Series, )
from openpyxl.chart.axis import DateAxis

ERI_dose_control = load_workbook('C:\\Work\\Python\\test2\\Для графика высвечивания.xlsx')
ws = ERI_dose_control.active

chart = ScatterChart()
chart.title = "Scatter Chart"
chart.style = 13
chart.x_axis.title = 'День с момента импульса'
chart.y_axis.title = 'Остаточная радиация, мкЗв/ч'

for type in range(0, 33):
    xvalues = Reference(ws, min_col=(type*3 + 1), min_row=2, max_row=11)
    values = Reference(ws, min_col=(type*3 + 2), min_row=1, max_row=11)
    graph = Series(values, xvalues, title_from_data=True)
    graph.marker.symbol = 'triangle'
    graph.marker.size = 5
    print(graph.marker.symbol)

    chart.series.append(graph)
ws.add_chart(chart, "A12")


ERI_dose_control.save('C:\\Work\\Python\\test2\\Для графика высвечивания.xlsx')
ERI_dose_control.close()


